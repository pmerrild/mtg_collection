from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path

import httpx
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app import create_app
from backend.domain import Inventory, parse_decklist, save_deck, text_export
from backend.importer import apply_import, parse_workbook, stage_import
from backend.scryfall import Scryfall
from backend.storage import ROOT, Store, now


HEADERS = ['Name', 'Type', 'Color', 'Rarity', 'Set#', 'Set', 'Count', 'Foil', 'Notes', 'Deck']


def write_book(path, rows):
    w = Workbook()
    s = w.active
    s.title = 'Input'
    s.append(HEADERS)
    for r in rows:
        s.append(r)
    w.save(path)
    w.close()


def source(name='Lightning Bolt', count=1, foil=None, number='001', set_code='abc', deck=None):
    return [name, 'Instant', 'Red', 'C', number, set_code, count, foil, None, deck]


def card(name='Lightning Bolt', number='1', set_code='abc', id='11111111-1111-1111-1111-111111111111', oracle='oracle-bolt', prices=None):
    return {'id': id, 'oracle_id': oracle, 'name': name, 'set': set_code, 'collector_number': number,
            'type_line': 'Instant', 'oracle_text': 'A test card.', 'color_identity': ['R'],
            'legalities': {'commander': 'legal', 'modern': 'legal'},
            'prices': prices or {'eur': '1.20', 'eur_foil': '2.30', 'usd': None, 'usd_foil': None}}


class VaultTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.store = Store(self.directory / 'data')
        self.book = self.directory / 'inventory.xlsx'

    def tearDown(self):
        self.temp.cleanup()

    def ingest(self, rows):
        write_book(self.book, rows)
        staged = stage_import(self.store, str(self.book))
        if not staged.get('unchanged'):
            apply_import(self.store, staged['id'])
        return staged

    def deck(self, text, name='Test deck', active=True, format='casual60'):
        return save_deck(self.store, {'name': name, 'format': format, 'active': active, 'decklist': text})

    def match(self, data, key=None):
        self.store.save_card(data)
        with self.store.connection() as db:
            db.execute('INSERT OR REPLACE INTO matches (printing_key,card_id,status,checked_at) VALUES (?, ?, ?, ?)',
                (key or f'{data["set"]}:{data["collector_number"]}', data['id'], 'matched', now()))

    @unittest.skipUnless(os.environ.get("MTG_ORIGINAL_WORKBOOK"), "Set MTG_ORIGINAL_WORKBOOK to run the original workbook acceptance check")
    def test_uploaded_workbook_counts_and_initial_decks(self):
        p = parse_workbook(os.environ['MTG_ORIGINAL_WORKBOOK'])
        self.assertEqual((p['rows'], p['copies']), (679, 892))
        self.assertEqual(sum(h['quantity'] for h in p['holdings'] if h['finish'] == 'foil'), 48)
        self.assertEqual(len(p['issues']), 1)
        self.assertEqual(p['issues'][0]['source_row'], 253)
        s = stage_import(self.store, os.environ['MTG_ORIGINAL_WORKBOOK'])
        apply_import(self.store, s['id'])
        inventory = Inventory(self.store)
        self.assertEqual(len(inventory.decks), 5)
        totals = sorted(inventory.deck(d['id'])['total'] for d in inventory.decks)
        self.assertEqual(totals, [19, 60, 60, 60, 60])

    def test_public_demo_workbook(self):
        result = parse_workbook(str(ROOT / 'sample-inventory.xlsx'))
        self.assertEqual((result['rows'], result['copies']), (1, 4))
        self.assertEqual(sum(h['quantity'] for h in result['holdings'] if h['finish'] == 'foil'), 1)
        self.assertEqual(result['issues'], [])

    def test_reimport_is_idempotent_and_keeps_row_ids(self):
        self.ingest([source(count=2, foil=1), source(count=3)])
        before = self.store.rows('SELECT * FROM holdings')
        self.assertTrue(stage_import(self.store, str(self.book))['unchanged'])
        self.assertEqual(before, self.store.rows('SELECT * FROM holdings'))
        self.assertEqual(sum(h['quantity'] for h in before), 5)

    def test_foil_split_conserves_quantity(self):
        result = self.ingest([source(count=4, foil=2)])
        collection = Inventory(self.store).collection()[0]
        self.assertEqual((collection['quantity'], collection['nonfoil'], collection['foil']), (4, 2, 2))

    def test_invalid_quantities_are_excluded_and_reported(self):
        result = self.ingest([source(count=None), source(count=2, foil=3), source(count=1.5), source(count=2)])
        self.assertEqual(result['copies'], 2)
        self.assertEqual(len(self.store.rows('SELECT * FROM issues')), 3)

    def test_failed_read_preserves_inventory(self):
        self.ingest([source(count=2)])
        before = self.store.rows('SELECT * FROM holdings')
        self.book.write_text('not a workbook')
        with self.assertRaises(ValueError):
            stage_import(self.store, str(self.book))
        self.assertEqual(before, self.store.rows('SELECT * FROM holdings'))

    def test_quantity_reduction_is_only_applied_after_review(self):
        self.ingest([source(count=3)])
        write_book(self.book, [source(count=1)])
        staged = stage_import(self.store, str(self.book))
        self.assertEqual(staged['reductions'][0]['before'], 3)
        self.assertEqual(sum(h['quantity'] for h in self.store.rows('SELECT * FROM holdings')), 3)
        apply_import(self.store, staged['id'])
        self.assertEqual(sum(h['quantity'] for h in self.store.rows('SELECT * FROM holdings')), 1)
        self.assertTrue(list((self.store.directory / 'backups').glob('*.sqlite3')))

    def test_collection_updates_preserve_edited_deck_targets(self):
        self.ingest([source(count=1, deck='Existing')])
        deck_id = self.store.rows('SELECT id FROM decks')[0]['id']
        save_deck(self.store, {'name': 'Existing', 'format': 'casual60', 'active': True, 'decklist': '4 Lightning Bolt\n1 Unowned Target'}, deck_id)
        self.ingest([source(count=2, deck='Existing')])
        self.assertEqual(Inventory(self.store).deck(deck_id)['total'], 5)

    def test_deleted_seeded_deck_is_not_recreated_by_future_imports(self):
        self.ingest([source(count=1, deck='Existing')])
        client = TestClient(create_app(self.store, background=False))
        deck = self.store.rows('SELECT id FROM decks')[0]['id']
        self.assertEqual(client.delete(f'/api/decks/{deck}').status_code, 200)
        self.ingest([source(count=2, deck='Existing')])
        self.assertEqual(self.store.rows('SELECT id FROM decks'), [])

    def test_same_card_in_two_decks_uses_physical_copy_once(self):
        self.ingest([source(count=1)])
        a = self.deck('1 Lightning Bolt', 'A')
        b = self.deck('1 Lightning Bolt', 'B')
        i = Inventory(self.store)
        self.assertEqual(i.wishlist({a, b}, 'assembled')['copies'], 1)
        self.assertEqual(i.wishlist({a, b}, 'shared')['copies'], 0)
        self.assertEqual(i.deck(b)['missing'], 0)
        self.assertEqual(i.deck(b)['missing_now'], 1)

    def test_draft_deck_does_not_reserve_copies(self):
        self.ingest([source()])
        self.deck('1 Lightning Bolt', 'Draft', active=False)
        active = self.deck('1 Lightning Bolt', 'Active')
        self.assertEqual(Inventory(self.store).deck(active)['missing_now'], 0)

    def test_sideboard_demand_is_not_double_covered(self):
        self.ingest([source(count=4)])
        deck = self.deck('Deck\n3 Lightning Bolt\nSideboard\n2 Lightning Bolt')
        self.assertEqual(Inventory(self.store).deck(deck)['missing'], 1)

    def test_exact_printing_is_allocated_before_flexible_requirement(self):
        self.ingest([source(number='1'), source(number='2')])
        flexible = self.deck('1 Lightning Bolt', 'Flexible')
        exact = self.deck('1 Lightning Bolt (ABC) 001', 'Exact')
        i = Inventory(self.store)
        self.assertEqual(i.wishlist({flexible, exact})['copies'], 0)
        self.assertEqual(i.deck(exact)['missing_now'], 0)

    def test_shared_mode_respects_exact_editions(self):
        self.ingest([source(number='1')])
        a = self.deck('1 Lightning Bolt (ABC) 1', 'A')
        b = self.deck('1 Lightning Bolt (ABC) 2', 'B')
        shared = Inventory(self.store).wishlist({a, b}, 'shared')
        self.assertEqual(shared['copies'], 1)
        self.assertEqual(shared['items'][0]['printing_key'], 'abc:2')

    def test_confirmed_printing_alias_matches_exact_target(self):
        self.ingest([source(number='001', set_code='old')])
        self.match(card(), 'old:1')
        deck = self.deck('1 Lightning Bolt (ABC) 1')
        self.assertEqual(Inventory(self.store).deck(deck)['missing'], 0)
        self.assertEqual(Inventory(self.store).wishlist({deck}, 'shared')['copies'], 0)

    def test_split_card_face_names_match_oracle_identity(self):
        self.ingest([source(name='Fire', number='1')])
        c = card(name='Fire // Ice', oracle='oracle-fire-ice')
        c['card_faces'] = [{'name': 'Fire'}, {'name': 'Ice'}]
        self.match(c)
        deck = self.deck('1 Fire // Ice')
        self.assertEqual(Inventory(self.store).deck(deck)['missing'], 0)

    def test_finish_correct_prices_and_unknown_currency_values(self):
        self.ingest([source(count=3, foil=1)])
        self.match(card())
        item = Inventory(self.store).collection()[0]
        self.assertEqual(item['value'], 4.7)
        self.assertEqual(item['priced_copies'], 3)
        self.store.set_settings({'currency': 'USD'})
        unknown = Inventory(self.store).collection()[0]
        self.assertIsNone(unknown['value'])
        self.assertEqual(unknown['priced_copies'], 0)

    def test_deck_parser_sections_printings_and_error_location(self):
        entries = parse_decklist('Commander\n1 Ramos, Dragon Engine\nDeck\n2x Lightning Bolt (ABC) 0001\nSideboard\n1 Lightning Bolt')
        self.assertEqual([e['zone'] for e in entries], ['commander', 'main', 'sideboard'])
        self.assertEqual(entries[1]['printing_key'], 'abc:1')
        with self.assertRaisesRegex(ValueError, 'Line 2'):
            parse_decklist('Deck\nnot a quantity')

    def test_text_export_aggregates_quantities_and_excludes_zero(self):
        self.assertEqual(text_export([{'name': 'Bolt', 'quantity': 2}, {'name': 'Bolt', 'quantity': 1}, {'name': 'Nothing', 'quantity': 0}]), '3 Bolt\n')

    def test_scryfall_batches_with_headers_and_caches_valid_matches(self):
        self.ingest([source()])
        requests = []
        def respond(request):
            requests.append(request)
            self.assertIn('MTGVault', request.headers['User-Agent'])
            self.assertEqual(request.headers['Accept'], 'application/json')
            if request.url.path == '/cards/collection':
                self.assertEqual(json.loads(request.content)['identifiers'], [{'set': 'abc', 'collector_number': '1'}])
                return httpx.Response(200, json={'data': [card()], 'not_found': []})
            raise AssertionError(request.url)
        client = Scryfall(self.store, transport=httpx.MockTransport(respond))
        client.refresh()
        self.assertEqual(self.store.rows('SELECT status FROM matches')[0]['status'], 'matched')
        self.assertEqual(Inventory(self.store).collection()[0]['value'], 1.2)
        client.refresh()
        self.assertEqual(len(requests), 1, 'Fresh matches should not cause another request')

    def test_scryfall_name_mismatch_requires_confirmation(self):
        self.ingest([source(name='Different Card')])
        client = Scryfall(self.store, transport=httpx.MockTransport(lambda request: httpx.Response(200, json={'data': [card()], 'not_found': []})))
        client.refresh()
        self.assertEqual(self.store.rows('SELECT status FROM matches')[0]['status'], 'mismatch')
        self.assertIsNone(Inventory(self.store).collection()[0]['value'])

    def test_api_failure_preserves_cached_prices(self):
        self.ingest([source()])
        self.match(card())
        def fail(request):
            raise httpx.ConnectError('offline', request=request)
        client = Scryfall(self.store, transport=httpx.MockTransport(fail))
        client.refresh(force=True)
        self.assertIn('unreachable', client.state['error'])
        self.assertEqual(Inventory(self.store).collection()[0]['value'], 1.2)

    def test_restart_preserves_decks(self):
        self.ingest([source()])
        deck = self.deck('1 Lightning Bolt\n1 Future Card')
        restarted = Store(self.store.directory)
        self.assertEqual(Inventory(restarted).deck(deck)['total'], 2)

    def test_api_exports_complete_targets_and_positive_shortages(self):
        self.ingest([source()])
        deck = self.deck('2 Lightning Bolt\n1 Future Card')
        client = TestClient(create_app(self.store, background=False))
        complete = client.get(f'/api/export?kind=deck&deck_id={deck}')
        self.assertIn('1 Future Card', complete.text)
        self.assertIn('2 Lightning Bolt', complete.text)
        missing = client.get(f'/api/export?kind=missing&deck_ids={deck}')
        self.assertIn('1 Lightning Bolt', missing.text)
        self.assertIn('1 Future Card', missing.text)
        state = client.get('/api/state').json()
        self.assertEqual(state['summary']['copies'], 1)

    def test_cross_origin_mutations_are_rejected(self):
        client = TestClient(create_app(self.store, background=False))
        r = client.patch('/api/settings', json={'currency': 'USD'}, headers={'Origin': 'https://unrelated.example'})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.store.settings()['currency'], 'EUR')

    def test_csv_escapes_formulas_in_source_notes(self):
        row = source()
        row[8] = '=HYPERLINK("bad")'
        self.ingest([row])
        # openpyxl writes leading '=' as a formula; model a literal Excel note.
        with self.store.connection() as db:
            db.execute('UPDATE holdings SET notes=?', ('=HYPERLINK("bad")',))
        client = TestClient(create_app(self.store, background=False))
        response = client.get('/api/export?kind=collection&format=csv')
        self.assertIn("'=HYPERLINK", response.text)

    def test_duplicate_column_headers_are_rejected(self):
        write_book(self.book, [source()])
        w = load_workbook(self.book)
        w['Input']['K1'] = 'Count'
        w.save(self.book)
        w.close()
        with self.assertRaisesRegex(ValueError, 'duplicate'):
            parse_workbook(str(self.book))

    def test_stale_preview_cannot_overwrite_a_newer_workbook(self):
        self.ingest([source(count=3)])
        write_book(self.book, [source(count=1)])
        staged = stage_import(self.store, str(self.book))
        write_book(self.book, [source(count=2)])
        with self.assertRaisesRegex(ValueError, 'changed since this preview'):
            apply_import(self.store, staged['id'])
        self.assertEqual(sum(h['quantity'] for h in self.store.rows('SELECT * FROM holdings')), 3)

    def test_numeric_zero_collector_number_is_preserved(self):
        self.ingest([source(number=0)])
        self.assertEqual(self.store.rows('SELECT printing_key FROM holdings')[0]['printing_key'], 'abc:0')

    def test_collection_requests_are_split_at_75_identifiers(self):
        self.ingest([source(name=f'Card {n}', number=str(n)) for n in range(1, 78)])
        sizes = []
        def respond(request):
            identifiers = json.loads(request.content)['identifiers']
            sizes.append(len(identifiers))
            result = [card(name=f'Card {x["collector_number"]}', number=x['collector_number'],
                           id=f'00000000-0000-0000-0000-{int(x["collector_number"]):012d}',
                           oracle=f'oracle-{x["collector_number"]}') for x in identifiers]
            return httpx.Response(200, json={'data': result, 'not_found': []})
        client = Scryfall(self.store, transport=httpx.MockTransport(respond))
        client.refresh()
        self.assertEqual(sizes, [75, 2])
        self.assertEqual(sum(c['match_status'] == 'matched' for c in Inventory(self.store).collection()), 77)


if __name__ == '__main__':
    unittest.main()
