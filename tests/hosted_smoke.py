"""Acceptance checks against a local Wrangler Worker. Uses only disposable local storage."""
import io,json,sys,urllib.request,urllib.error
from pathlib import Path
from openpyxl import load_workbook
BASE=sys.argv[1] if len(sys.argv)>1 else 'http://127.0.0.1:4173'
def request(path,method='GET',data=None,raw=False,headers=None):
 payload=data if raw else json.dumps(data).encode() if data is not None else None
 request_headers=dict(headers or ({'Content-Type':'application/json'} if payload else {}))
 if method not in ('GET','HEAD'): request_headers['X-Vault-Revision']=str(request('/state')['revision'])
 req=urllib.request.Request(BASE+'/api'+path,data=payload,method=method,headers=request_headers)
 with urllib.request.urlopen(req,timeout=20) as response:
  content=response.read()
  return json.loads(content) if response.headers.get('Content-Type','').startswith('application/json') else content
state=request('/state');assert state['summary']['copies']==892
original=Path(__file__).resolve().parents[1]/'MagicTheGatheringInventory.xlsx'
upload=lambda raw: request('/upload','POST',raw,True,{'Content-Type':'application/octet-stream','X-Workbook-Name':'acceptance.xlsx'})
assert upload(original.read_bytes())['unchanged']
try: upload(b'invalid workbook');raise AssertionError('Malformed workbook accepted')
except urllib.error.HTTPError as e: assert e.code==400
assert request('/state')['summary']['copies']==892
workbook=load_workbook(original);sheet=workbook['Input'];sheet.cell(2,7).value-=1
out=io.BytesIO();workbook.save(out)
pending=upload(out.getvalue());assert pending['reductions'];assert request('/state')['summary']['copies']==892
request('/import/'+pending['id']+'/apply','POST',{})
assert request('/state')['summary']['copies']==891
pending=upload(original.read_bytes());assert pending['pending'];assert request('/state')['summary']['copies']==891
request('/import/'+pending['id']+'/apply','POST',{});assert request('/state')['summary']['copies']==892
created=request('/decks','POST',{'name':'Hosted acceptance','format':'casual60','active':True,'decklist':'2 Unowned Acceptance Card\nSideboard\n1 Sol Ring'})
id=created['id'];state=request('/state');deck=next(d for d in state['decks'] if d['id']==id);assert deck['total']==3;assert '2 Unowned Acceptance Card' in deck['decklist']
assert request('/wishlist?deck_ids='+str(id))['copies']>=2
assert b'2 Unowned Acceptance Card' in request('/export?kind=deck&deck_id='+str(id))
request('/settings','PATCH',{'currency':'USD'});assert request('/state')['settings']['currency']=='USD'
request('/settings','PATCH',{'currency':'EUR'})
backup=request('/backup');assert backup['version']==1;assert backup['state']['holdings'];assert any(d['name']=='Hosted acceptance' for d in backup['state']['decks'])
assert sum(int(line.split(' ',1)[0]) for line in request('/export?kind=collection').decode().splitlines())==892
try: request('/settings','PATCH',{'currency':'USD'},headers={'Content-Type':'application/json','Origin':'https://example.com'});raise AssertionError('Foreign origin accepted')
except urllib.error.HTTPError as e: assert e.code==403
request('/decks/'+str(id),'DELETE');assert all(d['id']!=id for d in request('/state')['decks'])
print('Passed Worker storage, unchanged uploads, bad-file recovery, quantity review/apply, deck persistence, wishlist, exports, currency, backup, and same-origin checks.')
