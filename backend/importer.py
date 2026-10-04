from __future__ import annotations

import hashlib
import io
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .storage import Store, now


def collector(value) -> str:
    text = '' if value is None else str(value).strip()
    return str(int(text)) if text.isdigit() else text


def printing_key(set_code, number) -> str:
    return f'{str(set_code or "").strip().lower()}:{collector(number)}'


def integer(value, field: str, blank_zero=False) -> int:
    if value is None or value == '':
        if blank_zero:
            return 0
        raise ValueError(f'{field} is blank. Enter a quantity in Excel.')
    if isinstance(value, bool):
        raise ValueError(f'{field} must be a whole number.')
    try:
        n = float(value)
        if not n.is_integer() or n < 0:
            raise ValueError()
        return int(n)
    except (ValueError, TypeError, OverflowError):
        raise ValueError(f'{field} must be a nonnegative whole number.') from None


def parse_workbook(path: str, foil_mode='total') -> dict:
    if foil_mode not in {'total', 'additional', 'flag'}:
        raise ValueError('Choose the meaning of Count and Foil before importing.')
    file = Path(path).expanduser().resolve()
    if not file.is_file() or file.suffix.lower() != '.xlsx':
        raise ValueError('Choose an existing .xlsx workbook.')
    if file.stat().st_size > 32 * 1024 * 1024:
        raise ValueError('The workbook exceeds the 32 MB import limit.')
    before = file.stat()
    raw = file.read_bytes()
    after = file.stat()
    if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
        raise ValueError('Excel is still saving the workbook. Try Refresh again shortly.')
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ValueError('The workbook could not be read. Check the file and retry after Excel finishes saving.') from None
    try:
        if 'Input' not in workbook.sheetnames:
            raise ValueError('The workbook has no Input worksheet.')
        sheet = workbook['Input']
        if sheet.max_row > 100_000:
            raise ValueError('The Input sheet exceeds the 100,000-row import limit.')
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(x or '').strip() for x in next(iterator, [])]
        required = {'Name', 'Set', 'Set#', 'Count', 'Foil', 'Deck'}
        if len([h for h in headers if h]) != len(set(h for h in headers if h)):
            raise ValueError('Input contains duplicate column headings.')
        if not required.issubset(headers):
            raise ValueError('Input is missing columns: ' + ', '.join(sorted(required - set(headers))))
        holdings, issues = [], []
        populated = 0
        for row_number, cells in enumerate(iterator, 2):
            if not any(value is not None for value in cells):
                continue
            populated += 1
            row = dict(zip(headers, cells))
            name = str(row.get('Name') or '').strip()
            messages = []
            try:
                quantity = integer(row.get('Count'), 'Count')
                foils = integer(row.get('Foil'), 'Foil', blank_zero=True)
                if foil_mode == 'total':
                    if foils > quantity:
                        raise ValueError('Foil count exceeds Count. Check this row in Excel.')
                    quantities = {'nonfoil': quantity - foils, 'foil': foils}
                elif foil_mode == 'additional':
                    quantities = {'nonfoil': quantity, 'foil': foils}
                else:
                    if foils not in (0, 1):
                        raise ValueError('Foil flags must be blank, 0, or 1.')
                    quantities = {'foil' if foils else 'nonfoil': quantity}
            except ValueError as exc:
                quantities = {}
                messages.append(str(exc))
            set_value = row.get('Set')
            set_code = str(set_value or '').strip().lower()
            number = '' if row.get('Set#') is None else str(row['Set#']).strip()
            if not name:
                messages.append('Name is blank.')
                quantities = {}
            if not isinstance(set_value, str) or not re.fullmatch(r'[a-z0-9]{2,8}', set_code):
                messages.append('Set must be a confirmed Scryfall set code; no edition was guessed.')
            if not number:
                messages.append('Collector number is blank.')
            if messages:
                issues.append({'source_row': row_number, 'name': name or '(unnamed)',
                               'message': ' '.join(messages), 'raw': row})
            for finish, amount in quantities.items():
                if amount:
                    holdings.append({
                        'source_row': row_number, 'name': name,
                        'card_type': str(row.get('Type') or ''), 'color': str(row.get('Color') or ''),
                        'rarity': str(row.get('Rarity') or ''), 'set_code': set_code,
                        'collector_number': number, 'printing_key': printing_key(set_code, number),
                        'quantity': amount, 'finish': finish, 'notes': str(row.get('Notes') or ''),
                        'deck_label': str(row.get('Deck') or '').strip(),
                    })
        return {'fingerprint': hashlib.sha256(raw + foil_mode.encode()).hexdigest(),
                'source': str(file), 'rows': populated, 'holdings': holdings, 'issues': issues,
                'copies': sum(h['quantity'] for h in holdings), 'foil_mode': foil_mode}
    finally:
        workbook.close()


def snapshot_counts(holdings) -> Counter:
    counts = Counter()
    for h in holdings:
        counts[(h['printing_key'], h['finish'], h['name'].casefold())] += h['quantity']
    return counts


def stage_import(store: Store, path: str, foil_mode: str | None = None) -> dict:
    payload = parse_workbook(path, foil_mode or store.settings()['foil_mode'])
    with store.connection() as db:
        previous = db.execute("SELECT * FROM imports WHERE status='applied' ORDER BY id DESC LIMIT 1").fetchone()
        if previous and previous['fingerprint'] == payload['fingerprint']:
            return {'unchanged': True, 'id': previous['id'], 'rows': payload['rows'],
                    'copies': payload['copies'], 'issues': payload['issues'], 'reductions': []}
        old = snapshot_counts([dict(r) for r in db.execute('SELECT * FROM holdings')])
        current = snapshot_counts(payload['holdings'])
        labels = {(h['printing_key'], h['finish'], h['name'].casefold()): h['name'] for h in payload['holdings']}
        reductions = [{'name': labels.get(k, k[2]), 'printing_key': k[0], 'finish': k[1],
                       'before': v, 'after': current[k]} for k, v in old.items() if current[k] < v]
        payload['reductions'] = reductions
        pending = db.execute("SELECT id FROM imports WHERE status='pending' AND fingerprint=? ORDER BY id DESC LIMIT 1",
                             (payload['fingerprint'],)).fetchone()
        if pending:
            import_id = pending['id']
        else:
            import_id = db.execute('INSERT INTO imports (fingerprint,created_at,source,status,payload) VALUES (?,?,?,?,?)',
                (payload['fingerprint'], now(), payload['source'], 'pending', json.dumps(payload, default=str))).lastrowid
    return {'id': import_id, 'unchanged': False, 'rows': payload['rows'], 'copies': payload['copies'],
            'issues': payload['issues'], 'reductions': reductions}


def apply_import(store: Store, import_id: int) -> dict:
    with store.connection() as db:
        record = db.execute('SELECT * FROM imports WHERE id=?', (import_id,)).fetchone()
        if not record:
            raise ValueError('Import not found.')
        if record['status'] == 'applied':
            return {'id': import_id, 'unchanged': True}
        if record['status'] != 'pending':
            raise ValueError('This import preview was superseded. Refresh the workbook first.')
        latest = db.execute("SELECT id FROM imports WHERE status='pending' ORDER BY id DESC LIMIT 1").fetchone()
        if latest and latest['id'] != import_id:
            raise ValueError('A newer workbook preview exists. Review that preview first.')
    payload = json.loads(record['payload'])
    if parse_workbook(payload['source'], payload['foil_mode'])['fingerprint'] != record['fingerprint']:
        raise ValueError('The workbook changed since this preview. Refresh and review its latest contents first.')
    store.backup()
    with store.connection() as db:
        db.execute('DELETE FROM holdings')
        db.execute('DELETE FROM issues')
        for h in payload['holdings']:
            keys = list(h)
            db.execute(f'INSERT INTO holdings ({",".join(keys)}) VALUES ({",".join("?" for _ in keys)})', tuple(h.values()))
            db.execute("INSERT OR IGNORE INTO matches (printing_key,status) VALUES (?, 'unresolved')", (h['printing_key'],))
        for issue in payload['issues']:
            db.execute('INSERT INTO issues (source_row,name,message,raw) VALUES (?,?,?,?)',
                (issue['source_row'], issue['name'], issue['message'], json.dumps(issue['raw'], default=str)))
        decks = defaultdict(Counter)
        for h in payload['holdings']:
            if h['deck_label']:
                decks[h['deck_label']][h['name']] += h['quantity']
        saved_labels = db.execute("SELECT value FROM settings WHERE key='seeded_deck_labels'").fetchone()
        seeded_labels = set(json.loads(saved_labels['value'])) if saved_labels else set()
        for label, cards in decks.items():
            existing = db.execute('SELECT id FROM decks WHERE source_label=?', (label,)).fetchone()
            if existing or label in seeded_labels:
                continue  # Deck targets are app-owned and survive collection imports.
            deck_id = db.execute('INSERT INTO decks (name,format,source_label,seeded,updated_at) VALUES (?,?,?,?,?)',
                (label, 'commander' if 'commander' in label.casefold() else 'casual60', label, 1, now())).lastrowid
            for name, quantity in cards.items():
                db.execute('INSERT INTO entries (deck_id,name,quantity) VALUES (?,?,?)', (deck_id, name, quantity))
            seeded_labels.add(label)
        db.execute("UPDATE imports SET status='superseded' WHERE status='pending' AND id<>?", (import_id,))
        db.execute("UPDATE imports SET status='applied' WHERE id=?", (import_id,))
        for key, value in {'workbook_path': payload['source'], 'foil_mode': payload['foil_mode'],
                           'last_import_error': '', 'seeded_deck_labels': sorted(seeded_labels)}.items():
            db.execute('INSERT OR REPLACE INTO settings VALUES (?,?)', (key, json.dumps(value)))
    return {'id': import_id, 'copies': payload['copies'], 'rows': payload['rows'],
            'issues': len(payload['issues']), 'unchanged': False}
